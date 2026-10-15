"""Inventory SE 47/2026 PDFs and reconcile the supplied SIPASTI workbook.

Usage: python scripts/ahsp-inventory.py --corpus DIR --master FILE --output FILE
Requires openpyxl and pymupdf. Does not modify the workbook or the database.
"""
import argparse, collections, datetime, hashlib, json, pathlib, re, unicodedata, uuid
import openpyxl
import pymupdf

SOURCES = [('10892','Batang tubuh',''),('10894','Lampiran I',''),('10897','Lampiran II',''),('10899','Lampiran III','SMKK'),('10901','Lampiran IV','Sumber Daya Air'),('10903','Lampiran V','Bina Marga'),('10904','Lampiran VI','Cipta Karya / Perumahan'),('10906','Lampiran VII','')]
DIVISIONS = {
 'Cipta Karya / Perumahan': ['Persiapan lapangan / site work','Struktur','Arsitektur','Lansekap','Mekanikal dan elektrikal','Plambing','Jalan pada permukiman','Drainase jalan','Jaringan pipa di luar gedung','Sistem struktur RISHA','Tipologi RISHA'],
 'Bina Marga': ['Umum dan penerapan SMKK','Drainase','Tanah dan geosintetik','Preventif','Perkerasan berbutir dan beton semen','Perkerasan aspal','Struktur','Rehabilitasi jembatan','Pekerjaan harian dan lain-lain','Pemeliharaan'],
}

def clean(s): return re.sub(r'\s+',' ',str(s or '')).strip()
def token(s): return re.sub(r'[^\w]','',unicodedata.normalize('NFKC',clean(s)).lower())
def codekey(s): return re.sub(r'\s+','',clean(s)).lower()
def unitkey(s):
 s=token(s)
 return {'meterkubik':'m3','meterpersegi':'m2','meterpanjang':'m','m1':'m','bh':'buah','1buah':'buah','kilogram':'kg','orangjam':'oj','oranghari':'oh','lumsum':'ls','lumpsum':'ls'}.get(s,s)
def namekey(s):
 s=clean(s).lower().replace('∅','ø').replace('φ','ø')
 s=re.sub(r'\b1\s*(m[²³23]?|kg|buah|unit)\b','',s)
 return token(s)
def numeric(s):
 s=clean(s)
 if not re.fullmatch(r'\d+(?:[.,]\d+)?',s): return None
 return float(s.replace(',','.'))
def identity(*parts): return str(uuid.uuid5(uuid.NAMESPACE_URL,'ahsp:'+':'.join(parts)))
def quantity(name,unit=None):
 match=re.search(r"\b(\d+(?:[.,]\d+)?)\s*(m[²³23’']?|kg|buah|unit|batang|titik|set|ha|ton|tunggul)(?!\w)",name,re.I)
 if not match: return 1
 if unit and unitkey(match[2])!=unitkey(unit): return 1
 context=name[max(0,match.start()-30):match.start()]
 if re.search(r'(panjang|tinggi|diameter|kedalaman|jarak|ukuran|pjg|tiang|[TH]\s*=)\s*[-:]?\s*$',context,re.I) or re.match(r'^per[-\s]',name,re.I): return 1
 return numeric(match[1]) or 1

def divisions(field,code):
 if field in DIVISIONS:
  m=re.match(r'(\d+)',code)
  if m and 1<=int(m[1])<=len(DIVISIONS[field]): return f'Divisi {m[1]}: {DIVISIONS[field][int(m[1])-1]}'
 return ''

def headings(page,field):
 result=[]
 for block in page.get_text('dict')['blocks']:
  block_text=clean(' '.join(''.join(s['text'] for s in line['spans']) for line in block.get('lines',[])))
  for line in block.get('lines',[]):
   text=clean(''.join(s['text'] for s in line['spans']))
   if field=='Bina Marga':
    m=re.match(r'([A-J]\.\d+)\s+(.+)',text)
    if m: result.append({'code':m[1],'text':m[2],'y':line['bbox'][1]})
   else:
    m=re.match(r'((?:A\.\d+\.\S+|\d+(?:\.\d+){2,}[a-z]?))(?:\s+(.*))?$',text)
    if m:
     title=m[2] or ''
     if not title:
      whole=re.match(re.escape(m[1])+r'\s+(.+)',block_text)
      if whole and len(whole[1])<500 and not any(w in whole[1] for w in ['Koefisien','Uraian','Tenaga Kerja']): title=whole[1]
     result.append({'code':m[1].rstrip(':.'),'text':title,'y':line['bbox'][1]})
 return result

def component_table(table):
 rows=table.extract()
 header_index=next((i for i,r in enumerate(rows[:5]) if any('koefisien' in token(c) for c in r) and any('uraian' in token(c) for c in r)),None)
 if header_index is None: return None
 header=rows[header_index]
 def col(term): return next((i for i,c in enumerate(header) if term in token(c)),None)
 ni,ci,ui,ki=col('uraian'),col('kode'),col('satuan'),col('koefisien')
 if ni is None or ui is None or ki is None: return None
 group=None; components=[]; groups=set(); has_total=False; incomplete=False
 for row in rows[header_index+1:]:
  values=[clean(c) for c in row]; joined=' '.join(values)
  name=values[ni] if ni<len(values) else ''
  lead=token(joined)
  section=re.match(r'^([ABC])(?:\s|$)',values[0])
  heading=re.match(r'^(Tenaga\s+Kerja|Bahan|Peralatan|Alat)(?:\s|$)',name,re.I)
  if section:
   group={'A':'tenaga','B':'bahan','C':'alat'}[section[1]]; groups.add(group)
   if heading: name=re.sub(r'^\s*\*+\)?\s*','',name[heading.end():]).strip()
   if not name: continue
  elif token(name) in ('tenagakerja','bahan','peralatan','alat'):
   group={'tenagakerja':'tenaga','bahan':'bahan','peralatan':'alat','alat':'alat'}[token(name)]; groups.add(group); continue
  if ('hargasatuanpekerjaan' in lead or 'jumlahbiaya' in lead or 'jumlahharga' in lead) and ('d+e' in joined.lower() or values[0]=='F'): has_total=True
  if values[0] in ('D','E','F') or any(x in lead for x in ['jumlahharga','biayaumum','hargasatuanpekerjaan','catatan']):
   group=None if values[0] in ('D','E','F') else group
   continue
  if not group: continue
  if not name:
   if any(values[i] for i in (ui,ki) if i<len(values)) or (ci is not None and ci<len(values) and values[ci]): incomplete=True
   continue
  unit=values[ui] if ui<len(values) else ''
  coefficient=values[ki] if ki<len(values) else ''
  # Merged coefficient columns can place the actual value in an adjacent empty column.
  if not coefficient:
   for i in range(ki+1, min(ki+3,len(values))):
    if header[i] is None and values[i]: coefficient=values[i]; break
  if not unit:
   if name and not any(values[i] for i in range(ni+1,len(values))):
    if components: components[-1]['name']+=' '+name
    else: incomplete=True
    continue
   incomplete=True; continue
  # A merged PDF cell can contain several resources; do not invent their alignment.
  if (len(unit.split())>1 and unit.lower() not in ('orang hari','orang jam','orang bulan','unit hari')) or re.fullmatch(r'\d+(?:[.,]\d+)?(?:\s+\d+(?:[.,]\d+)?)+',coefficient):
   incomplete=True; continue
  if group=='tenaga' and unitkey(unit) not in ('oh','oj','ob'):
   incomplete=True; continue
  value=numeric(coefficient)
  if not coefficient: incomplete=True
  components.append({'id':str(len(components)+1),'group':group,'name':name,'code':values[ci] if ci is not None and ci<len(values) else '', 'unit':unit,'coefficient':value,'expression':coefficient if value is None else '', 'sourceLine':joined})
 if not components: return None
 status='extracted' if has_total and not incomplete and groups=={'tenaga','bahan','alat'} else 'needs_review'
 return {'components':components,'status':status,'bbox':table.bbox}

def inventory(corpus):
 items=[]; sources=[]; by_field_code=collections.defaultdict(list)
 for sid,annex,field in SOURCES:
  path=corpus/f'{sid}.pdf'; doc=pymupdf.open(path)
  sources.append({'id':sid,'title':f'SE 47/2026 - {annex}','annex':annex,'url':f'https://binakonstruksi.pu.go.id/?download_id={sid}&sdm_process_download=1','sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'page_count':len(doc),'checked_at':datetime.datetime.now(datetime.timezone.utc).isoformat(),'regulation':'SE 47/2026','assets':[]})
  if field not in ('Sumber Daya Air','Bina Marga','Cipta Karya / Perumahan'): continue
  bm_refs=collections.defaultdict(list); found=[]; analysis=collections.defaultdict(list)
  for pi,page in enumerate(doc):
   text=page.get_text(sort=True); hs=headings(page,field)
   if field=='Bina Marga' and pi>=30:
    for h in hs: bm_refs[h['code']].append(pi)
   catalogue=('normatif' in text.lower() and ('Kode' in text or 'KODE' in text) and ('Satuan' in text or 'SATUAN' in text))
   resource=('koefisien' in text.lower() and ('d+e' in text.lower() or 'a+b+c' in text.lower()))
   if not catalogue and not resource: continue
   for table in page.find_tables().tables:
    rows=table.extract()
    header=next((r for r in rows[:4] if any('kode'==token(c) for c in r) and any('normatif' in token(c) or 'tipeahsp' in token(c) for c in r)),None)
    if header:
     indices={token(c):i for i,c in enumerate(header) if c}
     codei=indices.get('kode'); namei=next((i for k,i in indices.items() if k in ('uraian','uraianpekerjaan','itempekerjaan')),None)
     uniti=indices.get('satuan'); typei=next((i for k,i in indices.items() if 'normatif' in k or k=='tipeahsp'),None); statusi=indices.get('status'); refi=indices.get('no')
     if None not in (codei,namei,uniti,typei):
      for row in rows:
       if len(row)<=max(codei,namei,uniti,typei): continue
       code=clean(row[codei]); name=clean(row[namei]); kind=clean(row[typei])
       if not code or token(code)=='kode' or not name or not any(t in kind.lower() for t in ['normatif','informatif']): continue
       found.append({'code':code,'name':name,'unit':clean(row[uniti]),'ahsp_type':kind,'regulation_status':clean(row[statusi]) if statusi is not None else '', 'catalogue_page':pi,'ref':clean(row[refi]) if refi is not None else ''})
    component=component_table(table) if resource else None
    if component:
     previous=[h for h in hs if h['y']<component['bbox'][1]]
     if previous: analysis[codekey(previous[-1]['code'])].append((pi,component))
   if pi%100==0: print('INVENTORY',sid,pi,'catalogue',len(found),'analyses',len(analysis),flush=True)
  # Duplicate source rows remain separate when their names or units differ.
  seen=set()
  for f in found:
   key=(codekey(f['code']),namekey(f['name']),unitkey(f['unit']))
   if key in seen: continue
   seen.add(key)
   candidates=analysis.get(codekey(f['code']),[])
   if field=='Bina Marga':
    refs=bm_refs.get(f['ref'],[])
    page=refs[0] if len(refs)==1 else f['catalogue_page']
    next_refs=sorted(p for v in bm_refs.values() for p in v if p>page)
    end=min(next_refs[0]-1 if next_refs else page+2,page+5,len(doc)-1) if refs else page
    component=None
   elif len(candidates)==1:
    page,component=candidates[0]; end=page
   else: page=f['catalogue_page']; end=page; component=None
   text='\n\n'.join(f'Halaman PDF {p+1}\n'+doc[p].get_text(sort=True) for p in range(page,end+1))
   printed=re.search(r'^\s*-?\s*(\d{3,5})\s*-?\s*$',doc[page].get_text(sort=True),re.M)
   overhead=re.search(r'(\d+)\s*%\s*[-–]\s*(\d+)\s*%',text)
   item={ 'id':identity(field,*key),'code':f['code'],'name':f['name'],'unit':f['unit'],'field':field,'division':divisions(field,f['code']),'regulation':'SE 47/2026','ahsp_type':f['ahsp_type'],'regulation_status':f['regulation_status'],'source_id':sid,'page_from':page+1,'page_to':end+1,'printed_page':printed[1] if printed else '', 'analysis_status':component['status'] if component else 'unavailable','components':component['components'] if component else [],'source_text':text,'overhead_min':float(overhead[1]) if overhead else None,'overhead_max':float(overhead[2]) if overhead else None,'output_quantity':quantity(f['name'],f['unit']), 'reconciliation_status':'hanya_regulasi','reconciliation_notes':[], 'master_rows':[] }
   if len(candidates)>1: item['reconciliation_notes'].append('Lebih dari satu tabel analisis memiliki kode yang sama; perlu memilih varian yang tepat.')
   if field=='Bina Marga': item['reconciliation_notes'].append('Daftar pekerjaan dapat dibaca; rincian koefisien dalam tabel gambar perlu OCR dan pemeriksaan terhadap lampiran.')
   items.append(item); by_field_code[(field,codekey(f['code']))].append(item)
  print('DONE',sid,'official',len(seen),'analysis tables',sum(len(v) for v in analysis.values()),flush=True)
 return items,sources,by_field_code

def reconcile(master,items,index):
 workbook=openpyxl.load_workbook(master,read_only=True,data_only=True)
 def records(name):
  sheet=workbook[name]; sheet.reset_dimensions(); rows=list(sheet.iter_rows(values_only=True)); return [dict(zip(rows[0],r)) for r in rows[1:] if r[0]]
 fields=[{k:clean(v) for k,v in r.items()} for r in records('Master_Fields')]
 raw=[]
 for ri,r in enumerate(records('Master_Items'),2):
  if r['regulasi']!='SE 47/2026' and not (r['regulasi']=='PERMEN 08/2023' and r['job_field_name']=='Umum'): continue
  candidates=index.get((clean(r['job_field_name']),codekey(r['code'])),[])
  notes=[]; item=None
  if r['regulasi']!='SE 47/2026': status='arsip'
  elif len(candidates)==1:
   item=candidates[0]
   if unitkey(r['unit'])!=unitkey(item['unit']): notes.append(f"Satuan master: {clean(r['unit'])}; satuan regulasi: {item['unit']}.")
   if namekey(r['name'])!=namekey(item['name']): notes.append(f"Nama master berbeda: {clean(r['name'])}")
   status='berbeda' if notes else 'cocok'
   item['master_rows'].append(ri)
   item['reconciliation_notes'].extend(notes)
   item['reconciliation_status']=status if item['reconciliation_status'] in ('hanya_regulasi','cocok') else item['reconciliation_status']
  elif len(candidates)>1:
   exact=[i for i in candidates if namekey(i['name'])==namekey(r['name']) and unitkey(i['unit'])==unitkey(r['unit'])]
   if len(exact)==1:
    item=exact[0]; status='cocok'; item['master_rows'].append(ri); item['reconciliation_status']=status
   else: status='perlu_ditinjau'; notes.append('Kode memiliki lebih dari satu varian di regulasi. Pemetaan otomatis tidak dipilih.')
  else: status='belum_ditemukan'; notes.append('Belum ditemukan padanan pada daftar resmi yang berhasil diekstrak. Status ini bukan bukti item tidak tercantum dalam regulasi.')
  if r['regulasi']=='SE 47/2026' and not item:
   item={ 'id':identity('master',str(ri),str(r['id'])),'code':clean(r['code']),'name':clean(r['name']),'unit':clean(r['unit']),'field':clean(r['job_field_name']),'division':divisions(clean(r['job_field_name']),clean(r['code'])) or clean(r['division_name']),'regulation':'SE 47/2026','ahsp_type':'','regulation_status':'','source_id':None,'page_from':None,'page_to':None,'printed_page':'','analysis_status':'unavailable','components':[],'source_text':'','overhead_min':None,'overhead_max':None,'output_quantity':1,'reconciliation_status':status,'reconciliation_notes':notes.copy(),'master_rows':[ri] }
   items.append(item)
  raw.append({'row_number':ri,'original_id':str(r['id']),'field_id':clean(r['job_field_id']),'regulation':clean(r['regulasi']),'code':clean(r['code']),'name':clean(r['name']),'unit':clean(r['unit']),'division':clean(r['division_name']),'raw':r,'item_id':item['id'] if item else None,'status':status,'notes':notes})
 for item in items:
  if len(item['master_rows'])>1:
   item['reconciliation_status']='perlu_ditinjau'; item['reconciliation_notes'].append('Lebih dari satu baris master dipetakan ke item ini.')
  item['reconciliation_notes']=list(dict.fromkeys(item['reconciliation_notes']))
 return fields,raw

if __name__=='__main__':
 parser=argparse.ArgumentParser(); parser.add_argument('--corpus',type=pathlib.Path,required=True); parser.add_argument('--master',type=pathlib.Path,required=True); parser.add_argument('--output',type=pathlib.Path,required=True); args=parser.parse_args()
 items,sources,index=inventory(args.corpus)
 fields,raw=reconcile(args.master,items,index)
 report={'items':len(items),'master_rows':len(raw),'current_master_rows':sum(r['regulation']=='SE 47/2026' for r in raw),'reconciliation':dict(collections.Counter(i['reconciliation_status'] for i in items)),'analysis':dict(collections.Counter(i['analysis_status'] for i in items)),'fields':dict(collections.Counter(i['field'] for i in items))}
 args.output.parent.mkdir(parents=True,exist_ok=True)
 args.output.write_text(json.dumps({'sources':sources,'fields':fields,'items':items,'master_rows':raw,'report':report},ensure_ascii=False,default=str),encoding='utf-8')
 print(json.dumps(report,ensure_ascii=True),flush=True)
